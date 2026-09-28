#!/usr/bin/env python3
"""Convert the MEXT chapter-2 workbook (sheet 表全体) to UTF-8 CSV.

Columns are found from the 成分識別子 row. Only 5-digit food codes are kept.
The 2026-03-27 errata workbook is compared to the data sheet. Cells that
still show the listed 誤 value are replaced with 正. Cells that already
match 正 are left as published.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from normalize import normalize_food_search_text

EDITION = "八訂増補2023"
ERRATA_VERSION = "2026-03-27"
SOURCE = "mext_sfct"
SOURCE_URL = "https://www.mext.go.jp/a_menu/syokuhinseibun/mext_00001.html"

TYPED_COLUMNS = (
    ("REFUSE", "refuse_pct"),
    ("ENERC_KCAL", "kcal"),
    ("PROT-", "protein_g"),
    ("PROTCAA", "protein_aa_g"),
    ("FAT-", "fat_g"),
    ("FATNLEA", "fat_tag_g"),
    ("CHOCDF-", "carb_g"),
    ("CHOAVLDF-", "carb_avail_g"),
    ("FIB-", "fiber_g"),
    ("NACL_EQ", "salt_eq_g"),
)

CSV_COLUMNS = [
    "food_code",
    "food_group",
    "index_no",
    "name",
    "display_name",
    "normalized_name",
    "reading",
    "base_amount",
    "unit_type",
    "refuse_pct",
    "kcal",
    "protein_g",
    "protein_aa_g",
    "fat_g",
    "fat_tag_g",
    "carb_g",
    "carb_avail_g",
    "fiber_g",
    "salt_eq_g",
    "raw_values",
    "estimated_fields",
    "source",
    "edition",
    "errata_version",
    "source_url",
]

# CI sample. Includes ごはん / 玄米 / 食パン / うどん / そば / 中華めん /
# 鶏卵 / 若どりむね / 牛乳 / 納豆 / 木綿豆腐, plus nearby names so ranking
# is not a single-row accident.
SAMPLE_FOOD_CODES = (
    "01004",
    "01026",
    "01034",
    "01039",
    "01048",
    "01064",
    "01085",
    "01088",
    "01089",
    "01111",
    "01128",
    "01137",
    "01154",
    "01198",
    "01200",
    "02017",
    "04032",
    "04033",
    "04046",
    "06061",
    "06182",
    "07012",
    "07107",
    "07148",
    "10003",
    "10134",
    "10144",
    "10253",
    "11123",
    "11220",
    "11288",
    "11227",
    "12004",
    "12005",
    "13003",
    "13025",
    "14017",
    "15023",
    "16037",
    "16045",
    "17042",
    "18001",
    "18031",
    "18057",
)

_ITEM_IDS = (
    ("差引き法による利用可能炭水化物", "CHOAVLDF-"),
    ("利用可能炭水化物（単糖当量）", "CHOAVLM"),
    ("利用可能炭水化物（質量計）", "CHOAVL"),
    ("レチノール活性当量", "VITA_RAE"),
    ("βクリプトキサンチン", "CRYPXB"),
    ("食物繊維総量", "FIB-"),
    ("炭水化物", "CHOCDF-"),
    ("食品名", "__NAME__"),
)


def _canon(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        text = format(value, "f").rstrip("0").rstrip(".")
        return text
    return str(value).strip()


def _compact(value: object) -> str:
    return re.sub(r"\s+", "", _canon(value))


def _sheet_rows(path: Path, sheet_name: str) -> list[list[object]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise SystemExit(f"{path} has no sheet {sheet_name!r}")
        sheet = workbook[sheet_name]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _find_identifier_row(rows: list[list[object]]) -> tuple[int, dict[str, int]]:
    for index, row in enumerate(rows):
        labels = {_compact(cell): col for col, cell in enumerate(row) if _compact(cell)}
        if "REFUSE" in labels and "ENERC_KCAL" in labels and "PROT-" in labels:
            return index, {key.strip(): col for key, col in (
                (str(cell).strip(), col)
                for col, cell in enumerate(row)
                if isinstance(cell, str) and cell.strip()
            )}
    raise SystemExit("成分識別子 row was not found")


def _locate_layout(rows: list[list[object]]) -> dict[str, int]:
    ident_index, ident_map = _find_identifier_row(rows)
    missing = [key for key, _ in TYPED_COLUMNS if key not in ident_map]
    if missing:
        raise SystemExit(f"missing component ids: {missing}")
    name_col = ident_map.get("成分識別子")
    if name_col is None:
        name_col = ident_map["REFUSE"] - 1
    if name_col < 3:
        raise SystemExit("food name column is not where the identifier row says")
    layout = {
        "identifier_row": ident_index,
        "group": name_col - 3,
        "code": name_col - 2,
        "index": name_col - 1,
        "name": name_col,
    }
    layout.update(ident_map)
    return layout


def _note_column(rows: list[list[object]], ident_index: int) -> int | None:
    for row in rows[: ident_index + 1]:
        for col, cell in enumerate(row):
            if isinstance(cell, str) and "備考" in cell:
                return col
    return None


def _food_code(value: object) -> str | None:
    if isinstance(value, int) and 0 <= value <= 99999:
        return f"{value:05d}"
    text = _canon(value)
    if re.fullmatch(r"\d{5}", text):
        return text
    if re.fullmatch(r"\d{1,5}", text):
        return text.zfill(5)
    return None


def _index_foods(rows: list[list[object]], code_col: int) -> dict[str, int]:
    found: dict[str, int] = {}
    for index, row in enumerate(rows):
        if code_col >= len(row):
            continue
        code = _food_code(row[code_col])
        if code:
            found[code] = index
    return found


# Underscores belong in ids such as ENERC_KCAL and NACL_EQ.
# A trailing hyphen belongs in ids such as PROT- and CHOCDF-.
_COMPONENT_ID = re.compile(r"^[A-Z][A-Z0-9]*(?:_[A-Z0-9]+)*-?$")


def _apply_name_fragment(
    rows: list[list[object]],
    foods: dict[str, int],
    code: str,
    col: int,
    wrong: object,
    right: object,
    stats: dict[str, int],
) -> None:
    """Chapter-2 name errata lists the changed fragment, not the whole cell."""
    row_index = foods.get(code)
    if row_index is None:
        stats["missing_food"] += 1
        return
    row = rows[row_index]
    text = "" if col >= len(row) or row[col] is None else str(row[col])
    wrong_text = "" if wrong is None else str(wrong)
    right_text = "" if right is None else str(right)
    if not wrong_text or wrong_text == right_text:
        return
    if wrong_text in text:
        row[col] = text.replace(wrong_text, right_text, 1)
        stats["patched"] += 1
        return
    if right_text and right_text in text:
        stats["already_reflected"] += 1
        return
    stats["unmatched"] += 1
    print(
        f"errata unmatched name {code}: cell={text!r} 誤={wrong_text!r} 正={right_text!r}",
        file=sys.stderr,
    )


def _apply_cell(
    rows: list[list[object]],
    foods: dict[str, int],
    code: str,
    col: int,
    wrong: object,
    right: object,
    stats: dict[str, int],
) -> None:
    if _canon(wrong) == _canon(right):
        return
    row_index = foods.get(code)
    if row_index is None:
        stats["missing_food"] += 1
        return
    row = rows[row_index]
    while len(row) <= col:
        row.append(None)
    current = row[col]
    if _canon(current) == _canon(right):
        stats["already_reflected"] += 1
        return
    if _canon(current) == _canon(wrong):
        row[col] = right
        stats["patched"] += 1
        return
    stats["unmatched"] += 1
    print(
        f"errata unmatched {code} col {col}: cell={current!r} 誤={wrong!r} 正={right!r}",
        file=sys.stderr,
    )


def apply_errata(
    rows: list[list[object]],
    layout: dict[str, int],
    errata_path: Path | None,
) -> dict[str, int]:
    stats = {
        "already_reflected": 0,
        "patched": 0,
        "unmatched": 0,
        "missing_food": 0,
        "skipped_item": 0,
    }
    if errata_path is None:
        print("errata workbook not provided; no comparison", file=sys.stderr)
        return stats

    foods = _index_foods(rows, layout["code"])
    errata = load_workbook(errata_path, data_only=True, read_only=True)
    try:
        if "本表" in errata.sheetnames:
            table_rows = list(errata["本表"].iter_rows(values_only=True))
            index = 0
            while index < len(table_rows):
                row = table_rows[index]
                if row and row[0] == "誤" and index + 1 < len(table_rows):
                    wrong_row = row
                    right_row = table_rows[index + 1]
                    code = _food_code(wrong_row[2] if len(wrong_row) > 2 else None)
                    if code:
                        width = max(len(wrong_row), len(right_row))
                        for errata_col in range(1, width):
                            wrong = wrong_row[errata_col] if errata_col < len(wrong_row) else None
                            right = right_row[errata_col] if errata_col < len(right_row) else None
                            _apply_cell(
                                rows,
                                foods,
                                code,
                                errata_col - 1,
                                wrong,
                                right,
                                stats,
                            )
                    index += 2
                    continue
                index += 1

        if "本表第2章" in errata.sheetnames:
            note_col = _note_column(rows, layout["identifier_row"])
            for row in errata["本表第2章"].iter_rows(values_only=True):
                if not row or len(row) < 8:
                    continue
                code = _food_code(row[2])
                item = _compact(row[5])
                if code is None or not item or item == "項目等":
                    continue
                if item == "各成分":
                    stats["skipped_item"] += 1
                    continue
                target = None
                for label, component_id in _ITEM_IDS:
                    if item == _compact(label):
                        target = component_id
                        break
                if target is None:
                    if "備考" in item and note_col is not None:
                        target_col = note_col
                    else:
                        stats["skipped_item"] += 1
                        continue
                elif target == "__NAME__":
                    _apply_name_fragment(
                        rows,
                        foods,
                        code,
                        layout["name"],
                        row[6],
                        row[7],
                        stats,
                    )
                    continue
                else:
                    target_col = layout.get(target)
                    if target_col is None:
                        stats["skipped_item"] += 1
                        continue
                _apply_cell(rows, foods, code, target_col, row[6], row[7], stats)
    finally:
        errata.close()

    print(
        "errata {version}: already_reflected={already_reflected} patched={patched} "
        "unmatched={unmatched} missing_food={missing_food} skipped_item={skipped_item}".format(
            version=ERRATA_VERSION, **stats
        ),
        file=sys.stderr,
    )
    return stats


def _raw_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return _canon(value)


def _parse_number(text: str) -> float | None:
    if text == "":
        return None
    try:
        return float(text)
    except ValueError:
        return None


def parse_component(value: object) -> tuple[float | None, bool, str]:
    """Return (number, estimated, original text).

    '(x)' -> x and estimated. '(0)' -> 0 and estimated.
    'Tr' and '(Tr)' -> 0, not estimated.
    '-', '*', '' -> null.
    """
    raw = _raw_text(value)
    if raw in ("", "-", "*"):
        return None, False, raw
    if raw in ("Tr", "(Tr)"):
        return 0.0, False, raw
    match = re.fullmatch(r"\((.+)\)", raw)
    if match:
        inner = match.group(1).strip()
        if inner == "Tr":
            return 0.0, False, raw
        number = _parse_number(inner)
        if number is None:
            return None, False, raw
        return number, True, raw
    return _parse_number(raw), False, raw


def _format_number(value: float | None) -> str:
    if value is None:
        return ""
    if float(value).is_integer():
        return str(int(value))
    return format(value, "f").rstrip("0").rstrip(".")


def convert_rows(
    rows: list[list[object]],
    layout: dict[str, int],
    only_codes: set[str] | None,
) -> list[dict[str, str]]:
    ident_map = {
        key: layout[key]
        for key in layout
        if _COMPONENT_ID.fullmatch(key)
    }
    output: list[dict[str, str]] = []
    seen: set[str] = set()
    for row in rows[layout["identifier_row"] + 1 :]:
        if layout["code"] >= len(row):
            continue
        code = _food_code(row[layout["code"]])
        if code is None or code in seen:
            continue
        if only_codes is not None and code not in only_codes:
            continue
        seen.add(code)
        name = row[layout["name"]] if layout["name"] < len(row) else ""
        name_text = name if isinstance(name, str) else _canon(name)
        raw_values: dict[str, str] = {}
        estimated: list[str] = []
        parsed: dict[str, float | None] = {}
        for component_id, col in ident_map.items():
            cell = row[col] if col < len(row) else None
            number, is_estimated, raw = parse_component(cell)
            raw_values[component_id] = raw
            if is_estimated:
                estimated.append(component_id)
            if any(component_id == key for key, _ in TYPED_COLUMNS):
                parsed[component_id] = number
        group = row[layout["group"]] if layout["group"] < len(row) else ""
        index_no = row[layout["index"]] if layout["index"] < len(row) else ""
        record = {
            "food_code": code,
            "food_group": _canon(group),
            "index_no": _canon(index_no),
            "name": name_text,
            "display_name": name_text,
            "normalized_name": normalize_food_search_text(name_text),
            "reading": "",
            "base_amount": "100",
            "unit_type": "g",
            "source": SOURCE,
            "edition": EDITION,
            "errata_version": ERRATA_VERSION,
            "source_url": SOURCE_URL,
            "raw_values": json.dumps(raw_values, ensure_ascii=False, sort_keys=True),
            "estimated_fields": json.dumps(estimated, ensure_ascii=False),
        }
        for component_id, column in TYPED_COLUMNS:
            record[column] = _format_number(parsed.get(component_id))
        output.append(record)
    output.sort(key=lambda item: item["food_code"])
    if only_codes is not None:
        missing = sorted(only_codes - {item["food_code"] for item in output})
        if missing:
            raise SystemExit(f"sample codes missing from workbook: {missing}")
    return output


def write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_COLUMNS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--errata", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument(
        "--sample",
        action="store_true",
        help="Write only the CI sample food codes",
    )
    parser.add_argument(
        "--codes",
        help="Comma-separated food codes to keep",
    )
    args = parser.parse_args(argv)

    rows = _sheet_rows(args.xlsx, "表全体")
    layout = _locate_layout(rows)
    stats = apply_errata(rows, layout, args.errata)
    only: set[str] | None = None
    if args.sample:
        only = set(SAMPLE_FOOD_CODES)
    elif args.codes:
        only = {code.zfill(5) for code in args.codes.split(",") if code.strip()}
    converted = convert_rows(rows, layout, only)
    write_csv(args.out, converted)
    print(
        f"wrote {len(converted)} foods to {args.out} "
        f"(errata patched={stats['patched']}, already_reflected={stats['already_reflected']})"
    )
    if len(converted) != 2538 and only is None:
        print(
            f"warning: expected 2538 five-digit foods, got {len(converted)}",
            file=sys.stderr,
        )
    return 0


if __name__ == "__main__":
    sys.exit(main())
